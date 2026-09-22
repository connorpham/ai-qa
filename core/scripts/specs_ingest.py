#!/usr/bin/env python3
"""specs_ingest.py — Ingest & convert external/binary design specs into Git-tracked Markdown.

Turns external specification files (e.g. Excel screen designs, PRD documents) into
version-tracked Markdown files inside `docs/specs/` with timestamps, author attribution,
and detailed update history (changelog).

Usage:
    python3 .ai-qa/scripts/specs_ingest.py --source <dir|file> [--out docs/specs]
    python3 .ai-qa/scripts/specs_ingest.py --audit [--specs docs/specs]
    python3 .ai-qa/scripts/specs_ingest.py --selftest

Python 3.9 compatible.
"""
import argparse
import os
import re
import sys
import shutil
import tempfile
from datetime import datetime

try:
    import openpyxl  # type: ignore
except ImportError:
    openpyxl = None


def clean_val(c):
    if c is None:
        return ""
    val = str(c).strip()
    return val.replace("\r\n", " ").replace("\n", " ").replace("|", "\\|")


def read_text(path):
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        return fh.read()


def spec_id_of(file_path):
    """The id a spec is cited by. From the filename, because that is the one
    thing a reader can see before opening anything."""
    fname = os.path.basename(file_path)
    m = re.search(r"(SC-[A-Z0-9]+-[A-Z0-9]+|[A-Z]{2,}-[A-Z0-9_-]+)", fname)
    return m.group(1) if m else os.path.splitext(fname)[0]


def parse_markdown_spec(file_path):
    """A spec already written in Markdown.

    It needs no conversion — only an identity, so the index can list it and a
    case can cite it. This function existed as the word `pass` for one release:
    the feature was announced, markdown sources were collected into nothing,
    and the run reported success.
    """
    text = read_text(file_path)
    mtime = datetime.fromtimestamp(os.path.getmtime(file_path)).strftime("%Y-%m-%d %H:%M:%S")
    title_m = re.search(r"(?m)^#\s+(.+?)\s*$", text)
    return {
        "id": spec_id_of(file_path),
        "title": title_m.group(1) if title_m else os.path.splitext(os.path.basename(file_path))[0],
        "file_mtime": mtime,
        "latest_version": "1.0",
        "latest_author": "",
        "latest_change": "Imported as written",
        "latest_date": mtime.split(" ")[0],
        "history": [],
        "functions": [],
        "validations": [],
        "source_file": file_path,
    }


def _yaml_str(value):
    """A YAML double-quoted scalar that survives a quote in a spreadsheet cell.

    Unescaped, one `"` in an author's name broke the frontmatter of the
    document the whole verification then cites.
    """
    return '"{}"'.format(str(value).replace("\\", "\\\\").replace('"', '\\"')
                         .replace("\n", " ").replace("\r", " "))


def render_md_frontmatter(data, rel_source):
    converted_at = datetime.now().strftime("%Y-%m-%dT%H:%M:%S")
    return "\n".join([
        "---",
        "id: " + _yaml_str(data["id"]),
        "title: " + _yaml_str(data["title"]),
        "version: " + _yaml_str(data["latest_version"]),
        "last_updated: " + _yaml_str(data["latest_date"]),
        "source_file: " + _yaml_str(rel_source),
        "source_mtime: " + _yaml_str(data["file_mtime"]),
        "converted_at: " + _yaml_str(converted_at),
        "---",
        "",
        "",
    ])


def parse_xlsx_spec(file_path):
    """Extract structured data from an Excel design specification file."""
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to parse .xlsx files. Run: pip install openpyxl")

    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    sheet_names = wb.sheetnames

    fname = os.path.basename(file_path)
    spec_id_m = re.search(r"(SC-[A-Z0-9]+-[A-Z0-9]+|[A-Z]{2,}-[A-Z0-9_-]+)", fname)
    spec_id = spec_id_m.group(1) if spec_id_m else os.path.splitext(fname)[0]

    title_m = re.search(r"(?:画面設計書|設計書|仕様書)[_-](.+?)\.xlsx", fname)
    title = title_m.group(1) if title_m else os.path.splitext(fname)[0]

    file_mtime = datetime.fromtimestamp(os.path.getmtime(file_path)).strftime("%Y-%m-%d %H:%M:%S")

    # 1. Update history / changelog
    history = []
    latest_version = "1.0"
    latest_author = ""
    latest_change = "Initial creation"
    latest_date = file_mtime

    for sname in ("更新履歴", "変更履歴", "History", "Changelog"):
        if sname in sheet_names:
            sheet = wb[sname]
            rows = list(sheet.iter_rows(values_only=True))
            start_idx = -1
            for idx, r in enumerate(rows):
                line_str = " ".join([str(c) for c in r if c is not None])
                if ("更新日" in line_str or "Date" in line_str) and ("内容" in line_str or "Description" in line_str or "Change" in line_str):
                    start_idx = idx + 1
                    break

            if start_idx != -1:
                for r in rows[start_idx:]:
                    vals = [clean_val(c) for c in r if c is not None]
                    if len(vals) >= 4 and vals[0] and vals[0].isdigit():
                        history.append({
                            "no": vals[0],
                            "date": str(vals[1]).split(" ")[0],
                            "author": vals[2],
                            "desc": vals[3]
                        })
                if history:
                    latest = history[-1]
                    latest_version = f"1.{len(history)}"
                    latest_author = latest["author"]
                    latest_change = latest["desc"]
                    latest_date = latest["date"]
            break

    # 2. Function overview / events
    functions = []
    for sname in ("機能概要", "Overview", "Functions"):
        if sname in sheet_names:
            sheet = wb[sname]
            rows = list(sheet.iter_rows(values_only=True))
            start_idx = -1
            for idx, r in enumerate(rows):
                line_str = " ".join([str(c) for c in r if c is not None])
                if "機能名" in line_str or "Function" in line_str:
                    start_idx = idx + 1
                    break

            if start_idx != -1:
                for r in rows[start_idx:]:
                    vals = [clean_val(c) for c in r if c is not None]
                    if len(vals) >= 3 and vals[0] and vals[0].isdigit():
                        f_table = vals[3] if len(vals) > 3 else "-"
                        functions.append({
                            "no": vals[0],
                            "name": vals[1],
                            "desc": vals[2],
                            "table": f_table
                        })
            break

    # 3. Validation / check rules
    validations = []
    for sname in ("チェック仕様", "Validation", "Checks"):
        if sname in sheet_names:
            sheet = wb[sname]
            rows = list(sheet.iter_rows(values_only=True))
            start_idx = -1
            for idx, r in enumerate(rows):
                line_str = " ".join([str(c) for c in r if c is not None])
                if "チェック仕様" in line_str or "Rule" in line_str or "Validation" in line_str:
                    start_idx = idx + 1
                    break

            if start_idx != -1:
                for r in rows[start_idx:]:
                    vals = [clean_val(c) for c in r if c is not None]
                    if len(vals) >= 3 and vals[0] and vals[0].isdigit():
                        v_type = vals[3] if len(vals) > 3 else "Error"
                        validations.append({
                            "no": vals[0],
                            "name": vals[1],
                            "desc": vals[2],
                            "type": v_type
                        })
            break

    wb.close()

    return {
        "id": spec_id,
        "title": title,
        "file_mtime": file_mtime,
        "latest_version": latest_version,
        "latest_author": latest_author,
        "latest_change": latest_change,
        "latest_date": latest_date,
        "history": history,
        "functions": functions,
        "validations": validations,
        "source_file": file_path,
    }


def render_spec_markdown(data, rel_source=None):
    """Render spec dictionary into standard Markdown with YAML frontmatter."""
    converted_at = datetime.now().strftime("%Y-%m-%dT%H:%M:%S")
    src = rel_source or os.path.basename(data["source_file"])

    lines = [
        "---",
        "id: " + _yaml_str(data["id"]),
        "title: " + _yaml_str(data["title"]),
        "version: " + _yaml_str(data["latest_version"]),
        "last_updated: " + _yaml_str(data["latest_date"]),
        "last_author: " + _yaml_str(data["latest_author"]),
        "last_change: " + _yaml_str(data["latest_change"]),
        "source_file: " + _yaml_str(src),
        "source_mtime: " + _yaml_str(data["file_mtime"]),
        "converted_at: " + _yaml_str(converted_at),
        "---",
        "",
        f"# [{data['id']}] {data['title']}",
        "",
        f"> **Source document:** `{src}`  ",
        f"> **Last modified:** `{data['file_mtime']}` · **Version:** `v{data['latest_version']}`  ",
        f"> **Converted at:** `{converted_at}`",
        "",
        "---",
        "",
        "## 1. Revision History (Changelog)",
        "",
        "| No | Date | Author | Change Description |",
        "|:---:|:---:|:---|:---|",
    ]

    if data["history"]:
        for h in data["history"]:
            lines.append(f"| {h['no']} | {h['date']} | {h['author']} | {h['desc']} |")
    else:
        lines.append(f"| 1 | {data['file_mtime']} | {data['latest_author'] or 'N/A'} | {data['latest_change']} |")

    lines.extend([
        "",
        "---",
        "",
        "## 2. Functions & Events (Overview)",
        "",
        "| No | Function Name | Description | Related Table / Entity |",
        "|:---:|:---|:---|:---|",
    ])

    if data["functions"]:
        for f in data["functions"]:
            lines.append(f"| {f['no']} | {f['name']} | {f['desc']} | `{f['table']}` |")
    else:
        lines.append("| - | - | Standard interface | - |")

    lines.extend([
        "",
        "---",
        "",
        "## 3. Validation & Business Rules",
        "",
        "| No | Rule Name | Condition & Expected Behavior | Type |",
        "|:---:|:---|:---|:---:|",
    ])

    if data["validations"]:
        for v in data["validations"]:
            lines.append(f"| {v['no']} | {v['name']} | {v['desc']} | `{v['type']}` |")
    else:
        lines.append("| - | - | Defined in form schema | - |")

    lines.append("")
    return "\n".join(lines)


def ingest_directory(source_dir, out_dir):
    """Scan and ingest all specification documents in a directory."""
    os.makedirs(out_dir, exist_ok=True)
    screens_dir = os.path.join(out_dir, "screens")
    os.makedirs(screens_dir, exist_ok=True)

    # A single file is a legal source. `--source <dir|file>` said so from the
    # first line of the usage text, and os.walk on a file yields nothing — so
    # the documented form reported "no specifications found" and exited 0.
    if os.path.isfile(source_dir):
        candidates = [(os.path.dirname(source_dir) or ".", [os.path.basename(source_dir)])]
        base = os.path.dirname(source_dir) or "."
    else:
        candidates = [(root, files) for root, _dirs, files in os.walk(source_dir)
                      if not any(p in ("[old]", "old", "bk", "BK", "対象外", "backup", "tmp")
                                 for p in root.split(os.sep))]
        base = source_dir

    active_files = []
    for root, files in candidates:
        for f in files:
            if f.startswith("~$"):
                continue
            # Markdown counts. The feature was announced, and the code said
            # `pass` — a spec already written in Markdown was silently dropped.
            if f.endswith(".xlsx") or (f.endswith(".md") and f not in ("README.md", "index.md")):
                active_files.append(os.path.join(root, f))

    if not active_files:
        # (converted, failed) — one shape for every exit, so a caller never has
        # to ask which kind of zero it just got back.
        print(f"No valid specification files found in {source_dir}", file=sys.stderr)
        return 0, 0

    index_entries = []
    count = 0

    failed = []
    for fpath in sorted(active_files):
        try:
            rel = os.path.relpath(fpath, base)
            if fpath.endswith(".md"):
                data = parse_markdown_spec(fpath)
                md_text = read_text(fpath)
                if not md_text.lstrip().startswith("---"):
                    md_text = render_md_frontmatter(data, rel) + md_text
            else:
                data = parse_xlsx_spec(fpath)
                md_text = render_spec_markdown(data, rel_source=rel)
            out_file = os.path.join(screens_dir, f"{data['id']}.md")

            with open(out_file, "w", encoding="utf-8") as fh:
                fh.write(md_text)

            index_entries.append({
                "id": data["id"],
                "title": data["title"],
                "category": os.path.dirname(rel) or "General",
                "file": f"screens/{data['id']}.md",
                "mtime": data["file_mtime"],
                "version": f"v{data['latest_version']}",
                "change": data["latest_change"]
            })
            count += 1
        except Exception as e:
            failed.append((os.path.basename(fpath), str(e)))
            print(f"Warning: Failed to convert {os.path.basename(fpath)}: {e}", file=sys.stderr)

    # Master index README.md
    index_lines = [
        "# Specifications Master Index",
        "",
        "> **Purpose:** Master catalog of digitized specifications with version and changelog tracking.",
        f"> **Total Specs:** `{count}` · **Generated at:** `{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}`",
        "",
        "---",
        "",
        "| Spec ID | Title | Category | Markdown File | Version | Source Modified | Latest Change |",
        "|:---|:---|:---|:---:|:---:|:---:|:---|",
    ]

    for it in index_entries:
        index_lines.append(f"| **`{it['id']}`** | {it['title']} | {it['category']} | [{it['id']}.md]({it['file']}) | `{it['version']}` | {it['mtime']} | {it['change']} |")

    index_path = os.path.join(out_dir, "README.md")
    with open(index_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(index_lines) + "\n")

    print(f"Converted {count} of {len(active_files)} specifications into {out_dir}")
    print(f"Master index generated at {index_path}")
    # A spec that did not convert is an oracle that is not there. Reporting that
    # as success is how a verification later cites a document nobody wrote.
    if failed:
        print(f"{len(failed)} source(s) did NOT convert — these specs do not exist to cite:",
              file=sys.stderr)
        for name, why in failed:
            print(f"  x {name}: {why}", file=sys.stderr)
    return count, len(failed)


def audit_specs(specs_dir):
    """What is in docs/specs, and what a case could not cite.

    `--audit` was in the usage text, in --help and in argparse, and `main()`
    never handled it: the flag printed the help and exited 1. A documented
    capability that silently does nothing is worse than a missing one, because
    somebody builds a process on top of it.
    """
    if not os.path.isdir(specs_dir):
        print("No specs directory at {}".format(specs_dir), file=sys.stderr)
        return 1

    found, problems = [], []
    for root, _dirs, files in os.walk(specs_dir):
        for f in sorted(files):
            if not f.endswith(".md") or f in ("README.md", "index.md"):
                continue
            p = os.path.join(root, f)
            rel = os.path.relpath(p, specs_dir)
            text = read_text(p)
            fm = re.match(r"(?s)^---\n(.*?)\n---\n", text)
            sid = ""
            if fm:
                m = re.search(r"(?m)^id:\s*\"?([^\"\n]+)\"?\s*$", fm.group(1))
                sid = (m.group(1).strip() if m else "")
            found.append((rel, sid))
            if not fm:
                problems.append((rel, "no YAML frontmatter — nothing records where this came from"))
            elif not sid:
                problems.append((rel, "frontmatter has no id — a case cannot cite it by name"))
            if len(text.strip()) < 40:
                problems.append((rel, "almost empty — it cannot decide what 'correct' means"))

    print("{} spec document(s) under {}".format(len(found), specs_dir))
    for rel, sid in found:
        print("  {:<48} {}".format(rel, sid or "(no id)"))
    if problems:
        print("\n{} problem(s) — these cannot serve as an oracle:".format(len(problems)),
              file=sys.stderr)
        for rel, why in problems:
            print("  x {}: {}".format(rel, why), file=sys.stderr)
        return 1
    return 0


def selftest():
    """Prove this converter can fail.

    The previous selftest built a dictionary by hand, rendered it, and asserted
    four substrings. Both parsers could be replaced with `return {}` and it
    still printed PASS — a gate that cannot go red, wired into `ai-qa doctor`,
    printing a tick beside work nobody had checked.

    So: mutations. Each one breaks something real, and each one must be caught.
    """
    tmp = tempfile.mkdtemp(prefix="aiqa-specs-test-")
    fails = []

    def expect(cond, msg):
        if not cond:
            fails.append(msg)

    try:
        src = os.path.join(tmp, "src")
        out = os.path.join(tmp, "out")
        os.makedirs(src)
        with open(os.path.join(src, "SC-ORD-001.md"), "w", encoding="utf-8") as fh:
            fh.write("# Order screen\n\n## 3.2 Totals\nR1 total = qty x price\n")

        # 1. a Markdown spec is INGESTED, not silently dropped
        count, failed = ingest_directory(src, out)
        expect(count == 1, "a markdown spec was not ingested (count={})".format(count))
        expect(failed == 0, "a clean run reported {} failures".format(failed))
        produced = os.path.join(out, "screens", "SC-ORD-001.md")
        expect(os.path.isfile(produced), "no output file for the markdown spec")
        body = read_text(produced) if os.path.isfile(produced) else ""
        expect(body.startswith("---"), "the ingested spec carries no frontmatter")
        expect("SC-ORD-001" in body, "the ingested spec lost its id")
        expect("R1 total = qty x price" in body, "the ingested spec lost the rule it exists to state")

        # 2. a single FILE is a legal source — the usage text always said so
        out2 = os.path.join(tmp, "out2")
        count2, _ = ingest_directory(os.path.join(src, "SC-ORD-001.md"), out2)
        expect(count2 == 1, "--source <file> ingested {} specs".format(count2))

        # 3. a source that cannot convert is REPORTED, not counted as success
        bad = os.path.join(tmp, "bad")
        os.makedirs(bad)
        with open(os.path.join(bad, "broken.xlsx"), "wb") as fh:
            fh.write(b"this is not a workbook")
        count3, failed3 = ingest_directory(bad, os.path.join(tmp, "out3"))
        expect(count3 == 0 and failed3 == 1,
               "a source that failed to convert was reported as {} converted / {} failed"
               .format(count3, failed3))

        # 4. an empty source is not a failure, and says so
        empty = os.path.join(tmp, "empty")
        os.makedirs(empty)
        expect(ingest_directory(empty, os.path.join(tmp, "out4")) == (0, 0),
               "an empty source directory did not report (0, 0)")

        # 5. a quote in a cell must not break the frontmatter of the document
        #    every later verdict cites
        hostile = {
            "id": 'SC-"X"-1', "title": 'The "Orders" screen', "file_mtime": "2026-09-22 10:00:00",
            "latest_version": "1.0", "latest_author": 'A "B" C', "latest_change": 'said "no"',
            "latest_date": "2026-09-22", "history": [], "functions": [], "validations": [],
            "source_file": "/tmp/x.xlsx",
        }
        md = render_spec_markdown(hostile, rel_source='a "b".xlsx')
        fm = re.match(r"(?s)^---\n(.*?)\n---\n", md)
        expect(fm is not None, "a quote in a cell destroyed the frontmatter block")
        if fm:
            for line in fm.group(1).splitlines():
                k, _, v = line.partition(":")
                v = v.strip()
                expect(v.startswith('"') and v.endswith('"') and len(v) >= 2,
                       "frontmatter field {!r} is not a quoted scalar: {!r}".format(k, v))

        # 6. --audit reports, and reports PROBLEMS rather than a clean bill
        expect(audit_specs(out) == 0, "--audit failed a directory it had just written")
        with open(os.path.join(out, "screens", "naked.md"), "w", encoding="utf-8") as fh:
            fh.write("# no frontmatter, nothing says where this came from\n" + "x" * 60 + "\n")
        expect(audit_specs(out) == 1, "--audit passed a spec with no frontmatter")
        expect(audit_specs(os.path.join(tmp, "nope")) == 1, "--audit passed a missing directory")

    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    if fails:
        print("specs_ingest --selftest FAILED", file=sys.stderr)
        for f in fails:
            print("  x {}".format(f), file=sys.stderr)
        return 1
    print("specs_ingest --selftest passed  (markdown ingested, single file, failures counted, "
          "quotes escaped, audit reports)")
    return 0


def main():
    parser = argparse.ArgumentParser(description="Ingest specifications into Git-tracked Markdown")
    parser.add_argument("--source", type=str, help="Source directory or file of specifications")
    parser.add_argument("--out", type=str, default="docs/specs", help="Target output directory (default: docs/specs)")
    parser.add_argument("--specs", type=str, default="docs/specs", help="Directory to audit (with --audit)")
    parser.add_argument("--audit", action="store_true", help="Audit existing specs and name the ones that cannot serve as an oracle")
    parser.add_argument("--selftest", action="store_true", help="Prove this converter can fail")

    args = parser.parse_args()

    if args.selftest:
        return selftest()
    if args.audit:
        return audit_specs(args.specs)
    if args.source:
        _count, failed = ingest_directory(args.source, args.out)
        # A spec that did not convert is an oracle that is not there, and an
        # exit code is the only part of this a script downstream can read.
        return 1 if failed else 0

    parser.print_help()
    return 1


if __name__ == "__main__":
    sys.exit(main() or 0)
